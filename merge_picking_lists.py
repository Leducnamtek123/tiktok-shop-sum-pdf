#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tool gộp và cộng dồn danh sách nhặt hàng (Picking List) TikTok Shop từ nhiều file PDF.
Tác giả: Hỗ trợ vận hành TikTok Shop.
"""

import os
import sys
import glob
import argparse
from datetime import datetime
from collections import defaultdict

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

try:
    import pdfplumber
except ImportError:
    print("Lỗi: Vui lòng cài đặt pdfplumber bằng lệnh: pip install pdfplumber")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

try:
    from rich.console import Console
    from rich.table import Table
    from rich.panel import Panel
    console = Console()
except ImportError:
    console = None

# Tọa độ các cột trong bảng PDF của TikTok Shop (đơn vị: points)
COL_NO = (15, 60)
COL_NAME = (120, 268)
COL_SKU = (268, 345)
COL_SELLER_SKU = (345, 405)
COL_QTY = (405, 445)
COL_ORDER_ID = (445, 585)


def parse_single_pdf(file_path):
    """
    Trích xuất toàn bộ dữ liệu từ 1 file Picking List PDF của TikTok Shop.
    Trả về:
      - metadata: dict {user, print_time, order_qty, product_qty, item_qty}
      - items: list of dict {no, product_name, sku, seller_sku, qty, order_ids}
    """
    filename = os.path.basename(file_path)
    metadata = {
        'filename': filename,
        'user': '',
        'print_time': '',
        'order_qty': 0,
        'product_qty': 0,
        'item_qty': 0
    }

    with pdfplumber.open(file_path) as pdf:
        # 1. Trích xuất metadata từ header trang 1
        p0_text = pdf.pages[0].extract_text() or ""
        for line in p0_text.split("\n"):
            line_str = line.strip()
            if line_str.startswith("User:"):
                metadata['user'] = line_str.replace("User:", "").strip()
            elif line_str.startswith("Print time:"):
                metadata['print_time'] = line_str.replace("Print time:", "").strip()
            elif "Order quantity:" in line_str:
                parts = line_str.split()
                for i, p in enumerate(parts):
                    if p == "quantity:" and i > 0:
                        prev = parts[i-1]
                        val = int(parts[i+1]) if i+1 < len(parts) and parts[i+1].isdigit() else 0
                        if prev == "Order":
                            metadata['order_qty'] = val
                        elif prev == "Product":
                            metadata['product_qty'] = val
                        elif prev == "Item":
                            metadata['item_qty'] = val

        # 2. Thu thập từ ngữ (words) theo từng trang
        all_skus = []
        current_product_no = None

        for p_idx, page in enumerate(pdf.pages):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            top_limit = 165 if p_idx == 0 else 55
            bottom_limit = 775
            body_words = [w for w in words if top_limit <= w['top'] <= bottom_limit]

            qty_words = [w for w in body_words if COL_QTY[0] <= w['x0'] <= COL_QTY[1] and w['text'].isdigit()]
            qty_words.sort(key=lambda w: w['top'])

            no_words = [w for w in body_words if COL_NO[0] <= w['x0'] <= COL_NO[1] and w['text'].isdigit()]
            no_words.sort(key=lambda w: w['top'])

            for q_idx, q in enumerate(qty_words):
                q_top = q['top']
                qty_val = int(q['text'])

                prev_nos = [n for n in no_words if n['top'] <= q_top + 10]
                if prev_nos:
                    associated_no = int(prev_nos[-1]['text'])
                else:
                    associated_no = current_product_no

                next_q_top = qty_words[q_idx + 1]['top'] - 2 if q_idx + 1 < len(qty_words) else 9999

                sku_words = [
                    w for w in body_words
                    if COL_SKU[0] <= w['x0'] <= COL_SKU[1]
                    and (q_top - 5) <= w['top'] < next_q_top
                ]
                sku_words.sort(key=lambda w: (round(w['top'], 1), w['x0']))
                sku_text = " ".join(w['text'] for w in sku_words).strip()

                seller_sku_words = [
                    w for w in body_words
                    if COL_SELLER_SKU[0] <= w['x0'] <= COL_SELLER_SKU[1]
                    and (q_top - 5) <= w['top'] < next_q_top
                ]
                seller_sku_text = " ".join(w['text'] for w in seller_sku_words).strip()

                order_words = [
                    w for w in body_words
                    if COL_ORDER_ID[0] <= w['x0'] <= COL_ORDER_ID[1]
                    and (q_top - 5) <= w['top'] < next_q_top
                    and w['text'].isdigit() and len(w['text']) >= 15
                ]
                order_ids = [w['text'] for w in order_words]

                all_skus.append({
                    'source_file': filename,
                    'page': p_idx + 1,
                    'no': associated_no,
                    'qty': qty_val,
                    'sku': sku_text,
                    'seller_sku': seller_sku_text,
                    'order_ids': order_ids,
                    'q_top': q_top
                })

                if associated_no is not None:
                    current_product_no = associated_no

        # 3. Thu thập tên sản phẩm theo số thứ tự (No)
        product_names = {}
        for p_idx, page in enumerate(pdf.pages):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            top_limit = 165 if p_idx == 0 else 55
            words = [w for w in words if top_limit <= w['top'] <= 775]

            no_words = [w for w in words if COL_NO[0] <= w['x0'] <= COL_NO[1] and w['text'].isdigit()]
            no_words.sort(key=lambda w: w['top'])

            name_words = [w for w in words if COL_NAME[0] <= w['x0'] <= COL_NAME[1]]

            if no_words:
                first_no_top = no_words[0]['top'] - 5
                before_first = [w for w in name_words if w['top'] < first_no_top]
                if before_first and current_product_no in product_names:
                    product_names[current_product_no].extend(before_first)

                for i, n in enumerate(no_words):
                    n_val = int(n['text'])
                    start_y = n['top'] - 5
                    end_y = no_words[i+1]['top'] - 5 if i + 1 < len(no_words) else 9999
                    in_range = [w for w in name_words if start_y <= w['top'] < end_y]
                    if n_val not in product_names:
                        product_names[n_val] = []
                    product_names[n_val].extend(in_range)
                    current_product_no = n_val
            else:
                if current_product_no in product_names:
                    product_names[current_product_no].extend(name_words)

        assembled_names = {}
        for n_val, w_list in product_names.items():
            w_list.sort(key=lambda w: (round(w['top'], 1), w['x0']))
            assembled_names[n_val] = " ".join(w['text'] for w in w_list).strip()

        for item in all_skus:
            item['product_name'] = assembled_names.get(item['no'], '')

        return metadata, all_skus


def merge_picking_lists(pdf_files):
    """
    Gộp và cộng dồn danh sách sản phẩm từ nhiều file PDF.
    """
    aggregated = defaultdict(lambda: {
        'product_name': '',
        'sku': '',
        'seller_sku': '',
        'total_qty': 0,
        'order_ids': set(),
        'by_file': defaultdict(int)
    })

    file_metadata_list = []
    all_order_ids_all_files = set()

    for fpath in pdf_files:
        meta, items = parse_single_pdf(fpath)
        file_metadata_list.append(meta)

        for it in items:
            key = (it['product_name'].strip(), it['sku'].strip(), it['seller_sku'].strip())
            entry = aggregated[key]
            entry['product_name'] = it['product_name'].strip()
            entry['sku'] = it['sku'].strip()
            entry['seller_sku'] = it['seller_sku'].strip()
            entry['total_qty'] += it['qty']
            entry['order_ids'].update(it['order_ids'])
            entry['by_file'][meta['filename']] += it['qty']
            all_order_ids_all_files.update(it['order_ids'])

    merged_list = []
    for key, data in aggregated.items():
        merged_list.append({
            'product_name': data['product_name'],
            'sku': data['sku'],
            'seller_sku': data['seller_sku'],
            'total_qty': data['total_qty'],
            'order_count': len(data['order_ids']),
            'order_ids': sorted(list(data['order_ids'])),
            'by_file': dict(data['by_file'])
        })

    merged_list.sort(key=lambda x: (-x['total_qty'], x['product_name']))

    total_stats = {
        'total_files': len(pdf_files),
        'file_names': [os.path.basename(f) for f in pdf_files],
        'total_orders_declared': sum(m['order_qty'] for m in file_metadata_list),
        'total_items_declared': sum(m['item_qty'] for m in file_metadata_list),
        'total_calculated_qty': sum(x['total_qty'] for x in merged_list),
        'total_unique_skus': len(merged_list),
        'total_unique_orders': len(all_order_ids_all_files)
    }

    return merged_list, total_stats, file_metadata_list


def export_to_excel(merged_list, total_stats, output_path):
    """
    Xuất file Excel tổng hợp được định dạng đẹp và chuyên nghiệp.
    """
    if not openpyxl:
        print("Cảnh báo: Cần cài openpyxl để xuất Excel (pip install openpyxl).")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tong Hop Picking List"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name='Segoe UI', size=16, bold=True, color='1F2937')
    font_sub = Font(name='Segoe UI', size=10, italic=True, color='4B5563')
    font_stat_lbl = Font(name='Segoe UI', size=10, bold=True, color='374151')
    font_stat_val = Font(name='Segoe UI', size=11, bold=True, color='1D4ED8')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_data = Font(name='Segoe UI', size=10)
    font_qty = Font(name='Segoe UI', size=11, bold=True, color='DC2626')
    font_total = Font(name='Segoe UI', size=11, bold=True, color='111827')

    fill_header = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_zebra = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
    fill_total = PatternFill(start_color='E0E7FF', end_color='E0E7FF', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    border_total = Border(
        top=Side(style='thin', color='1E3A8A'),
        bottom=Side(style='double', color='1E3A8A')
    )

    ws['A1'] = "BẢNG TỔNG HỢP DANH SÁCH NHẶT HÀNG (PICKING LIST TỔNG)"
    ws['A1'].font = font_title
    ws['A2'] = f"Ngày tổng hợp: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Nguồn: {total_stats['total_files']} file PDF TikTok Shop"
    ws['A2'].font = font_sub

    ws['A4'] = "Tổng số file PDF:"
    ws['B4'] = total_stats['total_files']
    ws['C4'] = "Tổng số đơn hàng:"
    ws['D4'] = total_stats['total_orders_declared']
    ws['E4'] = "Tổng số mặt hàng SKU:"
    ws['F4'] = total_stats['total_unique_skus']
    ws['G4'] = "TỔNG SỐ LƯỢNG LẤY:"
    ws['H4'] = total_stats['total_calculated_qty']

    for col in ['A', 'C', 'E', 'G']:
        ws[f'{col}4'].font = font_stat_lbl
    for col in ['B', 'D', 'F']:
        ws[f'{col}4'].font = font_stat_val
        ws[f'{col}4'].alignment = Alignment(horizontal='center')
    ws['H4'].font = Font(name='Segoe UI', size=13, bold=True, color='DC2626')
    ws['H4'].alignment = Alignment(horizontal='center')

    start_row = 6
    file_cols = total_stats['file_names']
    headers = ["STT", "Tên sản phẩm", "Phân loại (SKU)", "Mã Seller SKU", "TỔNG SỐ LƯỢNG", "Đã nhặt (Kho kiểm)"]
    if len(file_cols) > 1:
        headers.extend([f"SL ({f})" for f in file_cols])

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin
    ws.row_dimensions[start_row].height = 28

    current_row = start_row + 1
    for idx, item in enumerate(merged_list, 1):
        row_data = [
            idx,
            item['product_name'],
            item['sku'] if item['sku'] else "Mặc định",
            item['seller_sku'],
            item['total_qty'],
            "[   ]"
        ]
        if len(file_cols) > 1:
            for f in file_cols:
                row_data.append(item['by_file'].get(f, 0))

        is_even = (idx % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_qty if col_idx == 5 else font_data
            cell.border = border_thin
            if is_even:
                cell.fill = fill_zebra

            if col_idx == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [2, 3]:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif col_idx in [4, 6]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 5 or col_idx > 6:
                cell.alignment = Alignment(horizontal='right', vertical='center')

        ws.row_dimensions[current_row].height = 24
        current_row += 1

    ws.cell(row=current_row, column=1, value="")
    ws.cell(row=current_row, column=2, value="TỔNG CỘNG").font = font_total
    ws.cell(row=current_row, column=5, value=f"=SUM(E{start_row+1}:E{current_row-1})").font = font_qty
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = fill_total
        cell.border = border_total
        if c == 5:
            cell.alignment = Alignment(horizontal='right', vertical='center')

    col_widths = {1: 6, 2: 45, 3: 25, 4: 15, 5: 18, 6: 18}
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx in col_widths:
            ws.column_dimensions[col_letter].width = col_widths[col_idx]
        else:
            ws.column_dimensions[col_letter].width = 20

    wb.save(output_path)
    print(f"-> Đã xuất file Excel tổng hợp thành công: {output_path}")


def display_console_summary(merged_list, total_stats):
    """
    In kết quả tổng hợp ra màn hình console.
    """
    if console:
        console.print(Panel.fit(
            f"[bold green]KẾT QUẢ GỘP PICKING LIST TIKTOK SHOP[/bold green]\n"
            f"• Số file PDF đã gộp: [bold cyan]{total_stats['total_files']}[/bold cyan] file\n"
            f"• Tổng số đơn hàng: [bold cyan]{total_stats['total_orders_declared']}[/bold cyan] đơn\n"
            f"• Tổng số mặt hàng SKU khác nhau: [bold cyan]{total_stats['total_unique_skus']}[/bold cyan] SKU\n"
            f"• [bold red]TỔNG SỐ LƯỢNG SẢN PHẨM CẦN NHẶT: {total_stats['total_calculated_qty']}[/bold red] món",
            title="[bold blue]TikTok Shop Picking Consolidator[/bold blue]"
        ))

        table = Table(title="Danh sách hàng cần nhặt (Sắp xếp theo số lượng nhiều -> ít)", show_lines=True)
        table.add_column("STT", justify="center", style="dim", width=4)
        table.add_column("Tên sản phẩm", justify="left", style="white", min_width=35)
        table.add_column("Phân loại SKU", justify="left", style="yellow", width=22)
        table.add_column("Mã Seller", justify="center", style="dim", width=10)
        table.add_column("TỔNG SL", justify="right", style="bold red", width=10)

        for idx, it in enumerate(merged_list, 1):
            table.add_row(
                str(idx),
                it['product_name'][:45] + ("..." if len(it['product_name']) > 45 else ""),
                it['sku'] or "Mặc định",
                it['seller_sku'] or "-",
                str(it['total_qty'])
            )

        console.print(table)
    else:
        print("\n=== KẾT QUẢ GỘP PICKING LIST ===")
        print(f"Tổng files: {total_stats['total_files']} | Tổng SKU: {total_stats['total_unique_skus']} | Tổng SL cần nhặt: {total_stats['total_calculated_qty']}")
        for idx, it in enumerate(merged_list, 1):
            print(f"#{idx:2d} | SL: {it['total_qty']:3d} | SKU: {it['sku']:20s} | {it['product_name'][:40]}")


def main():
    parser = argparse.ArgumentParser(description="Gộp nhiều file Picking List PDF của TikTok Shop thành 1 bảng tổng hợp")
    parser.add_argument("files", nargs="*", help="Đường dẫn đến các file PDF hoặc thư mục chứa file PDF")
    parser.add_argument("--output", "-o", default="picking_list_tong_hop.xlsx", help="Tên file Excel đầu ra (mặc định: picking_list_tong_hop.xlsx)")
    parser.add_argument("--dir", "-d", help="Quét tất cả file PDF trong thư mục này")
    args = parser.parse_args()

    input_files = []
    if args.dir:
        input_files.extend(glob.glob(os.path.join(args.dir, "*Picking list*.pdf")))
        if not input_files:
            input_files.extend(glob.glob(os.path.join(args.dir, "*.pdf")))
    elif args.files:
        for item in args.files:
            if os.path.isdir(item):
                input_files.extend(glob.glob(os.path.join(item, "*Picking list*.pdf")))
            elif os.path.isfile(item):
                input_files.append(item)
            else:
                input_files.extend(glob.glob(item))
    else:
        default_dir = os.path.expanduser(r"~\Downloads")
        found = glob.glob(os.path.join(default_dir, "*Picking list*.pdf"))
        if found:
            print(f"Tìm thấy {len(found)} file Picking List trong Downloads.")
            input_files = found
        else:
            print("Vui lòng chỉ định đường dẫn file PDF hoặc thư mục chứa file PDF!")
            parser.print_help()
            sys.exit(1)

    input_files = sorted(list(set(input_files)))
    print(f"Đang xử lý {len(input_files)} file:")
    for f in input_files:
        print(f"  - {f}")

    merged_list, total_stats, _ = merge_picking_lists(input_files)
    display_console_summary(merged_list, total_stats)

    out_file = args.output
    export_to_excel(merged_list, total_stats, out_file)


if __name__ == "__main__":
    main()
